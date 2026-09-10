"""
집셀 매물 원본 엑셀에서 매매/전세/월세 전부(단, 매매는 dealType 필드 없이 기존 관례 유지) 추출.
전세/월세는 extract_jeonse_wolse.py와 같은 방식, 매매는 LISTINGS_ALL 기존 스키마와 동일하게 맞춤.

사용법: python3 extract_all_deals.py <집셀_파일들_있는_폴더> <출력.json>
"""
import openpyxl, glob, os, re, json, sys
from collections import defaultdict

# 갭투자 판단 키워드 — 대표 매물뿐 아니라 같은 매물을 올린 모든 중복 부동산의 "특징" 문구를 다 훑어서
# 하나라도 걸리면 갭투자로 집계. 세입자가 있어 즉시입주가 어려움을 뜻하는 표현들.
GAP_KEYWORDS = ['세안고', '전세안고', '갭투자', '긴잔금', '무주택', '긴전세', '임차인', '명도불가', '입주불가']


# 집셀 파일의 "단지 정보" 시트에 주소가 비어있는 날이 있어서(원인 불명, 가끔 발생),
# 예전에 정상적으로 확보했던 단지별 구/동 정보를 여기 하드코딩해두고 그날그날 대체용으로 씀.
# 새 단지가 계속 추가되니 이 딕셔너리는 주기적으로 최신 걸로 교체해줘야 함.
KNOWN_COMPLEX_LOCATION = {"e편한세상구성역플랫폼시티": ["기흥구", "마북동"], "강남마을5단지코오롱하늘채": ["기흥구", "구갈동"], "강남마을6단지자연앤": ["기흥구", "구갈동"], "기흥더샵프라임뷰": ["기흥구", "신갈동"], "기흥역더샵": ["기흥구", "구갈동"], "기흥역더퍼스트푸르지오": ["기흥구", "구갈동"], "기흥역롯데캐슬레이시티": ["기흥구", "구갈동"], "기흥역롯데캐슬스카이": ["기흥구", "신갈동"], "기흥역센트럴푸르지오": ["기흥구", "구갈동"], "기흥역파크푸르지오": ["기흥구", "구갈동"], "기흥푸르지오포레피스": ["기흥구", "영덕동"], "더샵보정애비뉴1단지": ["기흥구", "보정동"], "더샵보정애비뉴2단지": ["기흥구", "보정동"], "동백호수공원두산위브더제니스": ["기흥구", "동백동"], "블루밍구성더센트럴": ["기흥구", "마북동"], "삼거마을삼성래미안1차": ["기흥구", "마북동"], "상떼빌구성역플랫폼시티": ["기흥구", "보정동"], "성호샤인힐즈": ["기흥구", "보정동"], "신동백롯데캐슬에코1단지": ["기흥구", "중동"], "신동백롯데캐슬에코2단지": ["기흥구", "중동"], "신동백서해그랑블1차": ["기흥구", "중동"], "신동백서해그랑블2차": ["기흥구", "중동"], "신흥덕롯데캐슬레이시티": ["기흥구", "신갈동"], "연원마을삼성명가타운": ["기흥구", "보정동"], "연원마을엘지": ["기흥구", "마북동"], "용인기흥효성해링턴플레이스": ["기흥구", "영덕동"], "용인동백두산위브더제니스": ["기흥구", "동백동"], "용인보정꿈에그린": ["기흥구", "보정동"], "죽현마을동원로얄듀크": ["기흥구", "보정동"], "죽현마을아이파크": ["기흥구", "보정동"], "지석마을그대가크레던스": ["기흥구", "상하동"], "파크시엘": ["기흥구", "신갈동"], "행원마을동아솔레시티": ["기흥구", "보정동"], "힐스테이트기흥": ["기흥구", "구갈동"]}

# 원본 파일마다 같은 단지를 다른 이름으로 표기하는 경우가 있어서(집셀 vs 국토부, 혹은 부제 유무 등)
# 여기서 통일함. 안 그러면 실거래분석/시세현황에서 같은 단지가 두 개로 쪼개져서 잡힘.
COMPLEX_NAME_ALIASES = {
    '기흥역지웰푸르지오': '기흥역더퍼스트푸르지오',
    '기흥역롯데캐슬스카이(주상복합)': '기흥역롯데캐슬스카이',
}
def normalize_complex_name(name):
    return COMPLEX_NAME_ALIASES.get(name, name)


def detect_gap(items):
    for i in items:
        text = (i.get('feature') or '')
        # 원본에 공백이 섞여 들어오는 경우가 있어서(예: "긴 잔 금") 공백 제거하고 검사
        compact = text.replace(' ', '')
        if any(kw in compact for kw in GAP_KEYWORDS):
            return True
    return False

def parse_gu_dong(addr):
    parts = (addr or '').split()
    gu = parts[2] if len(parts) > 2 else '기흥구'
    # "동" 뒤에 지번(숫자)이 더 붙는 경우가 있어서(예: "구갈동 603") 마지막 토큰이 아니라
    # "동"으로 끝나는 토큰을 뒤에서부터 찾아야 함
    dong = ''
    for p in reversed(parts):
        if p.endswith('동'):
            dong = p
            break
    if not dong and parts:
        dong = parts[-1]
    return gu, dong

def process_file(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    info = {}
    if '단지 정보' in wb.sheetnames:
        ws = wb['단지 정보']
        for row in ws.iter_rows(min_row=1, max_row=999999, max_col=5, values_only=True):
            if row and row[0]:
                info[row[0]] = row[1]
    complex_name = normalize_complex_name(info.get('단지명') or '?')
    # "단지 정보" 시트 값이 키는 있는데 값만 비어있는(None) 날이 있어서, .get(키, 기본값)만 쓰면
    # 기본값이 아니라 None이 그대로 들어와버림 -- 키가 있어도 값이 None/빈 문자열이면 기본값을 쓰도록 or로 처리.
    addr = info.get('주소') or ''
    gu, dong = parse_gu_dong(addr)
    if not dong and complex_name in KNOWN_COMPLEX_LOCATION:
        # 오늘 파일에 주소가 비어있으면, 예전에 확보해둔 같은 단지의 구/동 정보로 대신 채움
        gu, dong = KNOWN_COMPLEX_LOCATION[complex_name]
    housing_type = info.get('부동산유형') or '아파트'
    if '오피스텔' in os.path.basename(path):
        housing_type = '오피스텔'

    ws = wb['매물 목록'] if '매물 목록' in wb.sheetnames else wb.active
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=999999, values_only=True):
        if not r or r[0] is None:
            continue
        gubun, rep_no, listing_no, deal_type, name, building, ho, floor, unit_type, area, price, rent, direction, checked, confirm, feature, agency, cp, dup, link, memo = (list(r) + [None]*21)[:21]
        rows.append({
            'gubun': gubun, 'rep_no': rep_no, 'listing_no': listing_no, 'dealType': deal_type,
            'building': str(building) if building else None, 'floor': floor, 'unitType': unit_type,
            'area': area, 'price': price, 'rent': rent, 'direction': direction,
            'checked': checked, 'agency': agency, 'feature': feature,
        })

    # 구분/대표매물번호로 그룹핑 (대표 매물 기준으로 중복 부동산 통합)
    groups = defaultdict(list)
    for r in rows:
        key = r['rep_no'] if r['gubun'] == '중복' else r['listing_no']
        groups[key].append(r)

    sale_out, jeonse_out, wolse_out = [], [], []
    for key, items in groups.items():
        rep = next((i for i in items if i['gubun'] == '대표'), items[0])
        dt = rep['dealType']
        agencies = list({i['agency'] for i in items if i['agency']})
        listing_ids = [i['listing_no'] for i in items if i['listing_no']]
        try:
            area_f = float(rep['area'])
        except:
            continue

        def normalize_price(raw):
            # 가격이 이미 억 단위(예: 11.8)인지 원단위(만원, 예: 118000)인지 방어적으로 처리
            if isinstance(raw, (int, float)) and raw > 1000:
                return raw / 10000
            return float(raw)

        if dt == '매매':
            # "대표" 매물 하나의 가격만 쓰면, 같은 매물을 올린 여러 부동산이 서로 다른 호가를 부른 경우
            # (예: 12억/12.5억 혼재) 대표로 뽑힌 쪽이 우연히 더 낮은 값이면 실제보다 낮게 보일 수 있음.
            # 그래서 이 그룹(대표+중복) 안에서 가장 높은 호가를 최종 가격으로 씀.
            try:
                prices = [normalize_price(i['price']) for i in items if i['price'] is not None]
                price_f = max(prices)
            except:
                continue
            sale_out.append({
                'gu': None, 'type': None, 'dong': None, 'complex': complex_name,
                'building': rep['building'], 'floor': str(rep['floor']), 'unitType': str(rep['unitType']),
                'area': area_f, 'price': round(float(price_f), 2), 'direction': rep['direction'],
                'checkedDate': str(rep['checked']), 'gap': detect_gap(items),
                'agencies': agencies, 'listingIds': listing_ids,
            })
        elif dt in ('전세', '월세'):
            try:
                deposits = [normalize_price(i['price']) for i in items if i['price'] is not None]
                deposit = max(deposits)
            except:
                continue
            entry = {
                'gu': None, 'type': None, 'dong': None, 'complex': complex_name,
                'building': rep['building'], 'floor': str(rep['floor']), 'unitType': str(rep['unitType']),
                'area': area_f, 'price': round(deposit, 2), 'direction': rep['direction'],
                'checkedDate': str(rep['checked']), 'dealType': dt,
                'agencies': agencies, 'listingIds': listing_ids,
            }
            if dt == '월세':
                try:
                    # 월세도 마찬가지로 중복 매물 중 최고 월세액을 사용
                    rents = [int(i['rent']) for i in items if i['rent']]
                    entry['rent'] = max(rents) if rents else None
                except:
                    entry['rent'] = None
                wolse_out.append(entry)
            else:
                jeonse_out.append(entry)

    for e in sale_out + jeonse_out + wolse_out:
        e['gu'] = gu; e['dong'] = dong; e['type'] = housing_type
    return sale_out, jeonse_out, wolse_out

def main():
    folder, out_path = sys.argv[1], sys.argv[2]
    files = sorted(glob.glob(os.path.join(folder, '집셀_*.xlsx')))
    all_sale, all_jeonse, all_wolse = [], [], []
    touched = set()
    for f in files:
        s, j, w = process_file(f)
        all_sale += s; all_jeonse += j; all_wolse += w
        if s or j or w:
            complex_name = (s+j+w)[0]['complex']
            housing_type = (s+j+w)[0]['type']
            touched.add((housing_type, complex_name))
    result = {'sale': all_sale, 'jeonse': all_jeonse, 'wolse': all_wolse, 'touched': sorted(list(touched))}
    json.dump(result, open(out_path, 'w', encoding='utf-8'), ensure_ascii=False)
    print(f"매매 {len(all_sale)}건 / 전세 {len(all_jeonse)}건 / 월세 {len(all_wolse)}건, 단지 {len(touched)}개 -> {out_path}")

if __name__ == '__main__':
    main()
